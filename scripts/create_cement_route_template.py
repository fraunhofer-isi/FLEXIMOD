# SPDX-FileCopyrightText: FLEXIMOD Developers
#
# SPDX-License-Identifier: AGPL-3.0-or-later

# ruff: noqa: E501

"""Create an Excel input template for every supported cement research route.

The workbook preserves the columns used by ``data/input/cement_plant_DE/plants.csv``
and adds the cement-model fields introduced for biomass, mixed RDF, oxyfuel and CCS.
It is an input-design aid: yellow cells require a route/fuel-specification decision;
green cells are model defaults or values copied from the reference plant.

Usage:
    python scripts/create_cement_route_template.py
    python scripts/create_cement_route_template.py --output C:/work/cement_routes.xlsx
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

REPO_ROOT = Path(__file__).resolve().parents[1]
REFERENCE_PLANTS = REPO_ROOT / "data" / "input" / "cement_plant_DE" / "plants.csv"
DEFAULT_OUTPUT = (
    REPO_ROOT / "data" / "input" / "cement_plant_DE" / "cement_route_parameter_template.xlsx"
)

MODEL_COLUMNS = [
    "route_id",
    "route_description",
    "name",
    "unit_type",
    "technology",
    "node",
    "objective",
    "fuel_type",
    "raw_meal_to_clinker_ratio",
    "waste_heat_per_t_clinker",
    "waste_heat_utilization_efficiency",
    "max_heat_out",
    "max_power",
    "min_heat_out",
    "min_power",
    "ramp_up",
    "ramp_down",
    "min_operating_steps",
    "min_down_steps",
    "initial_operational_status",
    "specific_heat_demand",
    "specific_electricity_aux",
    "eta_electric",
    "eta_fossil",
    "fossil_ng_share",
    "ng_co2_factor",
    "coal_co2_factor",
    "biomass_share",
    "rdf_share",
    "biomass_co2_factor",
    "biomass_co2_accounting_share",
    "rdf_mixed_fossil_co2_factor",
    "rdf_mixed_biogenic_co2_factor",
    "rdf_biogenic_co2_accounting_share",
    "calcination_emission_factor",
    "direct_separation_efficiency",
    "natural_gas_oxygen_demand",
    "coal_oxygen_demand",
    "biomass_oxygen_demand",
    "rdf_oxygen_demand",
    "hydrogen_oxygen_demand",
    "specific_oxygen_electricity_consumption",
    "capture_efficiency",
    "minimum_capture_fraction",
    "recovery_efficiency",
    "minimum_recovery_fraction",
    "specific_electricity_consumption",
    "specific_heat_consumption",
    "specific_variable_cost",
    "heat_cost",
    "storage_type",
    "max_capacity",
    "min_capacity",
    "max_power_charge",
    "max_power_discharge",
    "initial_soc",
    "storage_loss_rate",
    "efficiency_charge",
    "efficiency_discharge",
    "capacity",
    "min_soc",
    "max_soc",
    "efficiency",
    "oxygen_byproduct_t_per_mwh_hydrogen",
    "review_required",
]

CONVENTIONAL_FOSSIL_NG_SHARE = 0.034
R2_FOSSIL_NG_SHARE = 0.009 / (0.009 + 0.257)
R2_BIOMASS_SHARE = 0.245
R2_RDF_SHARE = 0.489
RDF_MIXED_FOSSIL_CO2_FACTOR = 0.243
RAW_MEAL_TO_CLINKER_RATIO = 1.55
WASTE_HEAT_PER_T_CLINKER = 0.22
WASTE_HEAT_UTILIZATION_EFFICIENCY = 0.90
BIOMASS_CO2_FACTOR = 0.403
RDF_MIXED_BIOGENIC_CO2_FACTOR = 0.135

AMINE_CAPTURE_EFFICIENCY = 0.90
AMINE_SPECIFIC_ELECTRICITY_CONSUMPTION = 0.12
AMINE_SPECIFIC_HEAT_CONSUMPTION = 0.95
AMINE_HEAT_COST = 0.0
CRYOGENIC_CAPTURE_EFFICIENCY = 0.90
CRYOGENIC_SPECIFIC_ELECTRICITY_CONSUMPTION = 0.11
OXYFUEL_RECOVERY_EFFICIENCY = 0.95
OXYFUEL_CCS_SPECIFIC_ELECTRICITY_CONSUMPTION = 0.10

NATURAL_GAS_OXYGEN_DEMAND = 0.147
COAL_OXYGEN_DEMAND = 0.248
HYDROGEN_OXYGEN_DEMAND = 0.24
SPECIFIC_OXYGEN_ELECTRICITY_CONSUMPTION = 0.22

ELECTROLYSER_EFFICIENCY = 0.709

ROUTES = (
    (
        "R1",
        "Conventional reference",
        (("preheater", "fossil"), ("simple_calciner", "fossil"), ("simple_kiln", "fossil")),
    ),
    (
        "R2",
        "Alternative fuel / waste-derived fuel",
        (("preheater", "fossil"), ("simple_calciner", "fossil"), ("simple_kiln", "fossil")),
    ),
    (
        "R3",
        "Conventional + amine CCS",
        (
            ("preheater", "fossil"),
            ("simple_calciner", "fossil"),
            ("simple_kiln", "fossil"),
            ("amine_ccs", ""),
        ),
    ),
    (
        "R4",
        "Conventional + cryogenic CCS",
        (
            ("preheater", "fossil"),
            ("simple_calciner", "fossil"),
            ("simple_kiln", "fossil"),
            ("cryogenic_ccs", ""),
        ),
    ),
    (
        "R5",
        "Electrified calciner",
        (("preheater", "fossil"), ("simple_calciner", "electricity"), ("simple_kiln", "fossil")),
    ),
    (
        "R6",
        "Hybrid-electric calciner",
        (
            ("preheater", "fossil"),
            ("simple_calciner", "hybrid_electricity_fossil"),
            ("simple_kiln", "fossil"),
        ),
    ),
    (
        "R7",
        "Electrified calciner + PCC",
        (
            ("preheater", "fossil"),
            ("simple_calciner", "electricity"),
            ("simple_kiln", "fossil"),
            ("amine_ccs", ""),
        ),
    ),
    (
        "R8",
        "Fully electrified clinker route + cryogenic CCS",
        (
            ("preheater", "electricity"),
            ("simple_calciner", "electricity"),
            ("simple_kiln", "electricity"),
            ("cryogenic_ccs", ""),
        ),
    ),
    (
        "R9",
        "Full oxyfuel",
        (
            ("preheater", "fossil"),
            ("oxyfuel_calciner", "fossil"),
            ("oxyfuel_kiln", "fossil"),
            ("oxyfuel_ccs", ""),
        ),
    ),
    (
        "R10",
        "Partial oxyfuel",
        (
            ("preheater", "fossil"),
            ("oxyfuel_calciner", "fossil"),
            ("simple_kiln", "fossil"),
            ("oxyfuel_ccs", ""),
        ),
    ),
    (
        "R11a",
        "Hydrogen-fired line, purchased hydrogen",
        (
            ("preheater", "fossil"),
            ("simple_calciner", "hydrogen"),
            ("simple_kiln", "hydrogen"),
            ("amine_ccs", ""),
        ),
    ),
    (
        "R11b",
        "Hydrogen-fired line, on-site hydrogen",
        (
            ("preheater", "fossil"),
            ("simple_calciner", "hydrogen"),
            ("simple_kiln", "hydrogen"),
            ("amine_ccs", ""),
            ("electrolyser", ""),
        ),
    ),
    (
        "R12a",
        "Oxy-hydrogen line, purchased hydrogen",
        (
            ("preheater", "fossil"),
            ("oxyfuel_calciner", "hydrogen"),
            ("oxyfuel_kiln", "hydrogen"),
            ("oxyfuel_ccs", ""),
        ),
    ),
    (
        "R12b",
        "Oxy-hydrogen line, on-site hydrogen",
        (
            ("preheater", "fossil"),
            ("oxyfuel_calciner", "hydrogen"),
            ("oxyfuel_kiln", "hydrogen"),
            ("oxyfuel_ccs", ""),
            ("electrolyser", ""),
        ),
    ),
    (
        "R13",
        "LEILAC + fossil heat",
        (("preheater", "fossil"), ("leilac_calciner", "fossil"), ("simple_kiln", "fossil")),
    ),
    (
        "R14",
        "Electrified LEILAC",
        (("preheater", "fossil"), ("leilac_calciner", "electricity"), ("simple_kiln", "fossil")),
    ),
    (
        "R15a",
        "Hydrogen LEILAC, purchased hydrogen",
        (("preheater", "fossil"), ("leilac_calciner", "hydrogen"), ("simple_kiln", "hydrogen")),
    ),
    (
        "R15b",
        "Hydrogen LEILAC, on-site hydrogen",
        (
            ("preheater", "fossil"),
            ("leilac_calciner", "hydrogen"),
            ("simple_kiln", "hydrogen"),
            ("electrolyser", ""),
        ),
    ),
    (
        "R16",
        "LEILAC + post-combustion CCS",
        (
            ("preheater", "fossil"),
            ("leilac_calciner", "fossil"),
            ("simple_kiln", "fossil"),
            ("amine_ccs", ""),
        ),
    ),
)

GREEN = PatternFill("solid", fgColor="E2F0D9")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
BLUE = PatternFill("solid", fgColor="D9EAF7")
GRAY = PatternFill("solid", fgColor="E7E6E6")
HEADER = PatternFill("solid", fgColor="1F4E78")


def read_reference_rows(path: Path) -> tuple[list[str], dict[str, dict[str, str]]]:
    with path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        rows = list(reader)
    by_technology = {row["technology"]: row for row in rows}
    return list(reader.fieldnames or []), by_technology


def empty_row() -> dict[str, object]:
    return {column: "" for column in MODEL_COLUMNS}


def reference_technology(technology: str) -> str | None:
    if technology in {"preheater", "simple_calciner", "simple_kiln", "thermal_storage"}:
        return technology
    if technology in {"leilac_calciner", "oxyfuel_calciner"}:
        return "simple_calciner"
    if technology == "oxyfuel_kiln":
        return "simple_kiln"
    return None


def stage_row(
    route_id: str,
    description: str,
    technology: str,
    fuel_type: str,
    reference_rows: dict[str, dict[str, str]],
) -> dict[str, object]:
    row = empty_row()
    reference_name = reference_technology(technology)
    if reference_name is not None:
        row.update(reference_rows.get(reference_name, {}))

    row.update(
        {
            "route_id": route_id,
            "route_description": description,
            "name": "<plant_name>",
            "unit_type": "cement_plant",
            "technology": technology,
            "node": "<node>",
            "objective": "min_variable_cost",
            "fuel_type": fuel_type,
            "raw_meal_to_clinker_ratio": RAW_MEAL_TO_CLINKER_RATIO,
            "waste_heat_per_t_clinker": WASTE_HEAT_PER_T_CLINKER,
            "waste_heat_utilization_efficiency": WASTE_HEAT_UTILIZATION_EFFICIENCY,
            "fossil_ng_share": CONVENTIONAL_FOSSIL_NG_SHARE,
            "biomass_share": 0.0,
            "rdf_share": 0.0,
            "biomass_co2_accounting_share": 0.0,
            "rdf_mixed_fossil_co2_factor": RDF_MIXED_FOSSIL_CO2_FACTOR,
            "rdf_biogenic_co2_accounting_share": 0.0,
            "initial_operational_status": 1,
        }
    )
    for field, requirement in {
        "max_heat_out": "<required MW_th>",
        "max_power": "<required MW_el>",
        "min_heat_out": "<required MW_th>",
        "min_power": "<required MW_el>",
        "ramp_up": "<required MW/step>",
        "ramp_down": "<required MW/step>",
    }.items():
        row[field] = requirement

    if route_id == "R2" and fuel_type in {"fossil", "hybrid_electricity_fossil"}:
        row.update(
            {
                "fossil_ng_share": R2_FOSSIL_NG_SHARE,
                "biomass_share": R2_BIOMASS_SHARE,
                "rdf_share": R2_RDF_SHARE,
                "biomass_co2_factor": BIOMASS_CO2_FACTOR,
                "rdf_mixed_biogenic_co2_factor": RDF_MIXED_BIOGENIC_CO2_FACTOR,
            }
        )

    if technology == "leilac_calciner":
        row["direct_separation_efficiency"] = 1.0
        row["review_required"] = "Confirm LEILAC direct separation efficiency."

    if technology in {"oxyfuel_calciner", "oxyfuel_kiln"}:
        used_fuels = {
            "fossil": {
                "natural_gas_oxygen_demand": NATURAL_GAS_OXYGEN_DEMAND,
                "coal_oxygen_demand": COAL_OXYGEN_DEMAND,
            },
            "hydrogen": {"hydrogen_oxygen_demand": HYDROGEN_OXYGEN_DEMAND},
            "hybrid_electricity_fossil": {
                "natural_gas_oxygen_demand": NATURAL_GAS_OXYGEN_DEMAND,
                "coal_oxygen_demand": COAL_OXYGEN_DEMAND,
            },
        }[fuel_type]
        row.update(used_fuels)
        row["specific_oxygen_electricity_consumption"] = SPECIFIC_OXYGEN_ELECTRICITY_CONSUMPTION

    return row


def ccs_row(route_id: str, description: str, technology: str) -> dict[str, object]:
    row = empty_row()
    row.update(
        {
            "route_id": route_id,
            "route_description": description,
            "name": "<plant_name>",
            "unit_type": "cement_plant",
            "technology": technology,
            "node": "<node>",
            "objective": "min_variable_cost",
            "specific_electricity_consumption": (
                AMINE_SPECIFIC_ELECTRICITY_CONSUMPTION
                if technology == "amine_ccs"
                else (
                    CRYOGENIC_SPECIFIC_ELECTRICITY_CONSUMPTION
                    if technology == "cryogenic_ccs"
                    else OXYFUEL_CCS_SPECIFIC_ELECTRICITY_CONSUMPTION
                )
            ),
            "specific_variable_cost": 0.0,
        }
    )
    if technology == "amine_ccs":
        row.update(
            {
                "capture_efficiency": AMINE_CAPTURE_EFFICIENCY,
                "minimum_capture_fraction": 0.0,
                "specific_heat_consumption": AMINE_SPECIFIC_HEAT_CONSUMPTION,
                "heat_cost": AMINE_HEAT_COST,
            }
        )
    elif technology == "cryogenic_ccs":
        row.update(
            {
                "capture_efficiency": CRYOGENIC_CAPTURE_EFFICIENCY,
                "minimum_capture_fraction": 0.0,
            }
        )
    else:
        row.update(
            {
                "recovery_efficiency": OXYFUEL_RECOVERY_EFFICIENCY,
                "minimum_recovery_fraction": 0.0,
            }
        )
    if technology == "amine_ccs":
        row["review_required"] = "heat_cost = 0 assumes internally supplied low-temperature heat."
    return row


def electrolyser_row(route_id: str, description: str) -> dict[str, object]:
    row = empty_row()
    row.update(
        {
            "route_id": route_id,
            "route_description": description,
            "name": "<plant_name>",
            "unit_type": "cement_plant",
            "technology": "electrolyser",
            "node": "<node>",
            "objective": "min_variable_cost",
            "max_power": "<required MW_el>",
            "min_power": 0.0,
            "efficiency": ELECTROLYSER_EFFICIENCY,
            "min_operating_steps": 1,
            "min_down_steps": 1,
            "initial_operational_status": 1,
            "oxygen_byproduct_t_per_mwh_hydrogen": 0.24,
            "review_required": "max_power must be sized from the plant scenario.",
        }
    )
    return row


def build_rows(reference_rows: dict[str, dict[str, str]]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for route_id, description, equipment in ROUTES:
        for technology, fuel_type in equipment:
            if technology in {"amine_ccs", "cryogenic_ccs", "oxyfuel_ccs"}:
                rows.append(ccs_row(route_id, description, technology))
            elif technology == "electrolyser":
                rows.append(electrolyser_row(route_id, description))
            else:
                rows.append(stage_row(route_id, description, technology, fuel_type, reference_rows))
    return rows


def optional_component_rows(reference_rows: dict[str, dict[str, str]]) -> list[dict[str, object]]:
    thermal = stage_row(
        "OPTIONAL", "Optional calciner thermal storage", "thermal_storage", "", reference_rows
    )
    thermal.update(
        {
            "name": "<plant_name>",
            "technology": "thermal_storage",
            "review_required": "Optional. Requires a calciner in the same plant.",
        }
    )
    hydrogen_store = empty_row()
    hydrogen_store.update(
        {
            "route_id": "OPTIONAL",
            "route_description": "Optional hydrogen buffer",
            "name": "<plant_name>",
            "unit_type": "cement_plant",
            "technology": "hydrogen_buffer_storage",
            "node": "<node>",
            "objective": "min_variable_cost",
            "capacity": "<required MWh_H2>",
            "min_soc": 0.0,
            "max_soc": 1.0,
            "initial_soc": 0.5,
            "efficiency_charge": 1.0,
            "efficiency_discharge": 1.0,
            "storage_loss_rate": 0.0,
            "review_required": "Optional. Requires an electrolyser in the same plant.",
        }
    )
    return [thermal, hydrogen_store]


def parameter_guide() -> list[tuple[str, str, str, str]]:
    return [
        (
            "route_id",
            "R2 selects alternative-fuel defaults; other routes use conventional defaults.",
            "Optional",
            "R2 or 2",
        ),
        (
            "raw_meal_to_clinker_ratio",
            "Raw meal required per tonne of clinker.",
            "Plant-level sizing and flow balance",
            "1.55",
        ),
        (
            "waste_heat_per_t_clinker",
            "Kiln waste heat available for the preheater.",
            "Plant-level sizing and flow balance",
            "0.22 MWh_th/t clinker",
        ),
        (
            "waste_heat_utilization_efficiency",
            "Usable fraction of available kiln waste heat.",
            "Plant-level sizing and flow balance",
            "0.90",
        ),
        (
            "max_heat_out",
            "Maximum stage heat output.",
            "Required for kiln-line stage",
            "No safe default",
        ),
        (
            "max_power",
            "Maximum stage electrical power.",
            "Required for kiln-line stage",
            "No safe default",
        ),
        (
            "min_heat_out",
            "Minimum stage heat output when operating.",
            "Required for kiln-line stage",
            "No safe default",
        ),
        (
            "min_power",
            "Minimum stage electrical power when operating.",
            "Required for kiln-line stage",
            "No safe default",
        ),
        (
            "ramp_up",
            "Maximum stage output increase per model step.",
            "Required for kiln-line stage",
            "No safe default",
        ),
        (
            "ramp_down",
            "Maximum stage output decrease per model step.",
            "Required for kiln-line stage",
            "No safe default",
        ),
        (
            "specific_heat_demand",
            "Heat required per tonne of stage output.",
            "Required for kiln-line stage",
            "Copied from reference where available",
        ),
        (
            "fuel_type",
            "electricity, fossil, hydrogen or hybrid_electricity_fossil.",
            "Required for kiln-line stage",
            "Set by route template",
        ),
        (
            "fossil_ng_share",
            "Natural-gas share inside the fossil remainder.",
            "Fossil/hybrid",
            "0.034 conventional; 0.0338345865 R2",
        ),
        (
            "biomass_share",
            "Separately procured biomass share of total combustion energy.",
            "Fossil/hybrid",
            "0.0 conventional; 0.245 R2",
        ),
        (
            "rdf_share",
            "Total mixed RDF share of total combustion energy.",
            "Fossil/hybrid",
            "0.0 conventional; 0.489 R2",
        ),
        (
            "biomass_co2_factor",
            "Physical biomass stack CO2 factor.",
            "Required when biomass_share > 0",
            "0.403 tCO2/MWh_th",
        ),
        (
            "rdf_mixed_fossil_co2_factor",
            "Fossil CO2 per MWh of total mixed RDF energy.",
            "RDF",
            "0.243 tCO2/MWh_th",
        ),
        (
            "rdf_mixed_biogenic_co2_factor",
            "Physical biogenic CO2 per MWh of total mixed RDF energy.",
            "Required when rdf_share > 0",
            "0.135 tCO2/MWh_th",
        ),
        (
            "*_co2_accounting_share",
            "Share of physical biogenic CO2 charged under the selected accounting case.",
            "Biomass/RDF",
            "0.0",
        ),
        (
            "direct_separation_efficiency",
            "LEILAC fraction of calcination CO2 directly separated.",
            "LEILAC",
            "1.0; confirm for study",
        ),
        (
            "*_oxygen_demand",
            "Oxygen need by fuel, per MWh_th fuel input.",
            "Oxyfuel",
            "NG 0.147; coal 0.248; hydrogen 0.24 tO2/MWh_th",
        ),
        (
            "specific_oxygen_electricity_consumption",
            "Electricity per tonne of oxygen generated internally.",
            "Oxyfuel",
            "0.22 MWh_el/tO2",
        ),
        (
            "capture_efficiency / recovery_efficiency",
            "Maximum physical capture/recovery fraction.",
            "CCS",
            "Amine/cryo 0.90; oxyfuel 0.95",
        ),
        (
            "specific_*_consumption",
            "CCS energy SEC.",
            "CCS",
            "Amine: 0.12 MWh_el/tCO2 and 0.95 MWh_th/tCO2; cryo: 0.11; oxyfuel: 0.10 MWh_el/tCO2",
        ),
        (
            "max_power, efficiency",
            "Electrolyser electrical rating and H2 conversion efficiency.",
            "On-site hydrogen",
            "Efficiency 0.709 MWh_H2/MWh_el; max_power has no safe default",
        ),
    ]


def add_table(sheet, name: str) -> None:
    if sheet.max_row < 2 or sheet.max_column < 1:
        return
    reference = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)


def style_sheet(sheet, *, freeze: str = "A2") -> None:
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = HEADER
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column_cells in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 36)
        sheet.column_dimensions[column_cells[0].column_letter].width = max(width, 12)
    sheet.row_dimensions[1].height = 32


def write_workbook(
    output: Path, reference_columns: list[str], reference_rows: dict[str, dict[str, str]]
) -> None:
    workbook = Workbook()
    readme = workbook.active
    readme.title = "Read me"
    readme.append(["Cement route parameter template"])
    readme.append(
        [
            "Purpose",
            "Create one plants.csv per route by filtering plants_template on route_id and replacing yellow placeholders.",
        ]
    )
    readme.append(["Green", "Model default or value copied from cement_plant_DE/plants.csv."])
    readme.append(
        [
            "Yellow",
            "A required route-, technology- or fuel-specification value. Replace before solving.",
        ]
    )
    readme.append(["Blue", "Route structure or identifier supplied by this template."])
    readme.append(
        [
            "Important",
            "RDF means total mixed RDF energy. Use rdf_share, never rdf_excl_biomass_share.",
        ]
    )
    readme.append(
        [
            "Important",
            "RDF fossil factor is 0.243 tCO2/MWh_th of total mixed RDF. Do not multiply it by a fossil fraction again.",
        ]
    )
    readme.append(
        [
            "Important",
            "The template's preheater fuel choices are modelling assumptions inherited from the current reference case; review them for your source data.",
        ]
    )
    for row in readme.iter_rows(min_row=1, max_row=1):
        for cell in row:
            cell.fill = HEADER
            cell.font = Font(color="FFFFFF", bold=True, size=14)
    readme.column_dimensions["A"].width = 18
    readme.column_dimensions["B"].width = 110
    for row in readme.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index in range(2, readme.max_row + 1):
        readme.row_dimensions[index].height = 34

    matrix = workbook.create_sheet("Route matrix")
    matrix.append(
        [
            "route_id",
            "description",
            "technologies",
            "external hydrogen",
            "electrolyser",
            "CCS boundary note",
        ]
    )
    for route_id, description, equipment in ROUTES:
        technologies = " -> ".join(technology for technology, _ in equipment)
        has_electrolyser = any(technology == "electrolyser" for technology, _ in equipment)
        has_hydrogen = any(fuel_type == "hydrogen" for _, fuel_type in equipment)
        note = "Oxyfuel CCS receives only oxyfuel-stage CO2." if route_id == "R10" else ""
        matrix.append(
            [
                route_id,
                description,
                technologies,
                "yes" if has_hydrogen and not has_electrolyser else "no",
                "yes" if has_electrolyser else "no",
                note,
            ]
        )
    style_sheet(matrix)
    add_table(matrix, "CementRouteMatrix")

    template = workbook.create_sheet("plants_template")
    template.append(MODEL_COLUMNS)
    for row in build_rows(reference_rows):
        template.append([row[column] for column in MODEL_COLUMNS])
    style_sheet(template, freeze="A2")
    add_table(template, "CementPlantsTemplate")
    for row in template.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("<required"):
                cell.fill = YELLOW
            elif cell.column <= 8:
                cell.fill = BLUE
            elif cell.value not in ("", None):
                cell.fill = GREEN
    review_column = get_column_letter(MODEL_COLUMNS.index("review_required") + 1)
    template.conditional_formatting.add(
        f"A2:{get_column_letter(template.max_column)}{template.max_row}",
        FormulaRule(formula=[f"NOT(ISBLANK(${review_column}2))"], fill=YELLOW),
    )

    optional = workbook.create_sheet("optional_components")
    optional.append(MODEL_COLUMNS)
    for row in optional_component_rows(reference_rows):
        optional.append([row[column] for column in MODEL_COLUMNS])
    style_sheet(optional)
    add_table(optional, "CementOptionalComponents")
    for row in optional.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("<required"):
                cell.fill = YELLOW
            elif cell.value not in ("", None):
                cell.fill = GREEN

    guide = workbook.create_sheet("parameter_guide")
    guide.append(["field", "meaning", "when required", "default / source"])
    for row in parameter_guide():
        guide.append(row)
    style_sheet(guide)
    add_table(guide, "CementParameterGuide")

    reference = workbook.create_sheet("reference_plants_csv")
    reference.append(reference_columns)
    for row in reference_rows.values():
        reference.append([row.get(column, "") for column in reference_columns])
    style_sheet(reference)
    add_table(reference, "CementReferencePlants")

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--reference", type=Path, default=REFERENCE_PLANTS)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    if not args.reference.exists():
        raise FileNotFoundError(f"Reference cement plants.csv not found: {args.reference}")
    columns, rows = read_reference_rows(args.reference)
    write_workbook(args.output, columns, rows)
    print(f"Wrote cement route template: {args.output}")


if __name__ == "__main__":
    main()
